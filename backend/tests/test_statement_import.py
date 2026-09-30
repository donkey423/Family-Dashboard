from datetime import date, datetime
from contextlib import closing
from pathlib import Path
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from family_finance_hub.config import Settings
from family_finance_hub.documents.processors.pdf import ProcessedPdf
from family_finance_hub.finance.statements.contracts import (
    ReconciliationSummary,
    StatementData,
    StatementLine,
    StatementParseResult,
)
from family_finance_hub.main import create_app
from family_finance_hub.models import FinanceTransaction, Statement
from family_finance_hub.finance.statements.taiwan_credit_cards import TaiwanCreditCardStatementParser


class FakePdfProcessor:
    def process(self, request, *, password_candidates=()):
        return ProcessedPdf(
            preview_bytes=request.content,
            extracted_text="synthetic statement text",
            page_count=2,
            was_encrypted=False,
            ocr_required=False,
            ocr_status="not_required",
            successful_candidate_index=None,
        )


class FakeStatementParser:
    bank_id = "synthetic-bank"
    format_version = "synthetic-v1"

    def parse(self, extracted_text: str, *, page_count: int) -> StatementParseResult:
        assert extracted_text == "synthetic statement text"
        assert page_count == 2
        return StatementParseResult(
            status="parsed",
            statement=StatementData(
                bank_id=self.bank_id,
                format_version=self.format_version,
                period_start=date(2026, 9, 1),
                period_end=date(2026, 9, 30),
                account_hint="****1234",
                lines=(
                    StatementLine(
                        line_index=1,
                        transaction_date=date(2026, 9, 3),
                        posting_date=date(2026, 9, 4),
                        description="超市",
                        transaction_kind="purchase",
                        amount=-120,
                        currency="TWD",
                    ),
                    StatementLine(
                        line_index=2,
                        transaction_date=date(2026, 9, 8),
                        description="退貨",
                        transaction_kind="refund",
                        amount=20,
                        currency="TWD",
                    ),
                    StatementLine(
                        line_index=3,
                        transaction_date=date(2026, 9, 15),
                        description="自動繳款",
                        transaction_kind="payment",
                        amount=300,
                        currency="TWD",
                    ),
                ),
                reconciliation=ReconciliationSummary(status="matched", basis="synthetic", difference=0),
            ),
        )


class VerifiedCardPdfProcessor:
    def process(self, request, *, password_candidates=()):
        return ProcessedPdf(
            preview_bytes=request.content,
            extracted_text=(
                "115/09/06 0 37 0 0 0 +37\n"
                "115/08/16 115/08/17\n37 8224 TW\neToro\n"
                "TWQR\n(02)2745-8080"
            ),
            page_count=2,
            was_encrypted=False,
            ocr_required=False,
            ocr_status="not_required",
            successful_candidate_index=None,
        )


def make_client(tmp_path: Path, parser=None):
    settings = Settings(
        database_url=f"sqlite:///{(tmp_path / 'statement.db').as_posix()}",
        storage_root=tmp_path / "documents",
        excel_output_path=tmp_path / "家庭收支記錄.xlsx",
    )
    app = create_app(
        settings,
        create_schema=True,
        pdf_processor=FakePdfProcessor(),
        statement_parser=parser,
    )
    return TestClient(app), settings


def upload_pdf(client: TestClient) -> str:
    response = client.post(
        "/api/documents",
        files={"file": ("synthetic-statement.pdf", b"synthetic pdf bytes", "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()["id"]


def test_analysis_without_parser_is_pending_and_does_not_create_transactions(tmp_path):
    client, _settings = make_client(tmp_path)
    with client:
        document_id = upload_pdf(client)
        response = client.post(f"/api/documents/{document_id}/statement-analysis", json={})
        assert response.status_code == 200
        assert response.json()["status"] == "pending"
        assert response.json()["reason_code"] == "parser_not_configured"
        assert client.get("/api/finance/transactions").json()["total"] == 0
        statement_id = response.json()["statement_id"]
        confirm = client.post(
            f"/api/statements/{statement_id}/confirm",
            json={"review_version": response.json()["review_version"]},
        )
        assert confirm.status_code == 409


def test_approved_interest_policy_is_reviewable_and_preserved_through_excel(tmp_path):
    class InterestPolicyParser(FakeStatementParser):
        bank_id = "cathay"

        def parse(self, extracted_text, *, page_count):
            result = super().parse(extracted_text, page_count=page_count)
            undated = StatementLine(
                line_index=4, posting_date=date(2026, 9, 30), description="循環利息",
                transaction_kind="interest", amount=-5, currency="TWD",
            )
            return result.model_copy(update={"statement": result.statement.model_copy(update={
                "lines": (*result.statement.lines, undated), "closing_date": date(2026, 9, 30),
            })})

    client, settings = make_client(tmp_path, InterestPolicyParser())
    with client:
        document_id = upload_pdf(client)
        analyzed = client.post(f"/api/documents/{document_id}/statement-analysis", json={}).json()
        assert analyzed["status"] == "ready"
        detail = client.get(f"/api/statements/{analyzed['statement_id']}").json()
        interest = detail["lines"][-1]
        assert interest["transaction_date"] is None
        assert interest["effective_transaction_date"] == "2026-09-30"
        assert interest["transaction_date_basis"] == "statement_closing_date"
        assert client.get("/api/finance/transactions").json()["total"] == 0
        confirm = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": analyzed["review_version"]})
        assert confirm.status_code == 200
        repeated = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": analyzed["review_version"]})
        assert repeated.json()["reused"]
        assert client.get("/api/finance/transactions").json()["total"] == 4
        with client.app.state.session_factory() as session:
            transaction = session.scalar(select(FinanceTransaction).where(FinanceTransaction.transaction_kind == "interest"))
            assert transaction.transaction_date == date(2026, 9, 30)
            raw = json.loads(transaction.raw_json)
            assert raw["source_transaction_date"] is None
            assert raw["transaction_date_basis"] == "statement_closing_date"
            assert session.get(Statement, analyzed["statement_id"]).statement_json["lines"][-1]["transaction_date"] is None
        assert client.post("/api/exports/excel/settings", json={"enabled": True}).json()["exported_transactions"] == 4
        with closing(load_workbook(settings.workbook_path, data_only=True)) as workbook:
            row = next(workbook["信用卡月支出"].iter_rows(min_row=2, values_only=True))
            assert row == (datetime(2026, 9, 1), "國泰世華信用卡", "TWD", 120, 20, 100, 0, 5, 105, 300, 4)


def test_statement_analysis_confirmation_is_idempotent_and_projects_card_sheet(tmp_path):
    client, settings = make_client(tmp_path, FakeStatementParser())
    with client:
        account = client.post(
            "/api/statement-accounts",
            json={"bank_id": "synthetic-bank", "display_name": "測試卡", "account_hint": "****1234"},
        )
        assert account.status_code == 200
        document_id = upload_pdf(client)

        analysis = client.post(
            f"/api/documents/{document_id}/statement-analysis",
            json={"statement_account_id": account.json()["id"]},
        )
        assert analysis.status_code == 200
        analyzed = analysis.json()
        assert analyzed["status"] == "ready"
        assert analyzed["line_count"] == 3
        assert "extracted_text" not in analyzed

        detail = client.get(f"/api/statements/{analyzed['statement_id']}")
        assert detail.status_code == 200
        assert len(detail.json()["lines"]) == 3
        assert "extracted_text" not in detail.json()

        confirmation = client.post(
            f"/api/statements/{analyzed['statement_id']}/confirm",
            json={"review_version": analyzed["review_version"]},
        )
        assert confirmation.status_code == 200
        assert confirmation.json()["transaction_count"] == 3
        assert confirmation.json()["reused"] is False

        repeated = client.post(
            f"/api/statements/{analyzed['statement_id']}/confirm",
            json={"review_version": analyzed["review_version"]},
        )
        assert repeated.status_code == 200
        assert repeated.json()["reused"] is True
        assert client.get("/api/finance/transactions").json()["total"] == 3

        export = client.post("/api/exports/excel/settings", json={"enabled": True})
        assert export.status_code == 200
        assert export.json()["exported_transactions"] == 3
        with closing(load_workbook(settings.workbook_path, data_only=True)) as workbook:
            rows = list(workbook["信用卡月支出"].iter_rows(min_row=2, values_only=True))
            assert rows == [(datetime(2026, 9, 1), "測試卡", "TWD", 120, 20, 100, 0, 0, 100, 300, 3)]


def test_verified_bank_parser_auto_creates_local_account_and_imports(tmp_path):
    settings = Settings(
        database_url=f"sqlite:///{(tmp_path / 'verified.db').as_posix()}",
        storage_root=tmp_path / "documents",
        excel_output_path=tmp_path / "家庭收支記錄.xlsx",
    )
    app = create_app(
        settings,
        create_schema=True,
        pdf_processor=VerifiedCardPdfProcessor(),
        statement_parser=TaiwanCreditCardStatementParser(),
    )
    client = TestClient(app)
    with client:
        document_id = upload_pdf(client)
        analysis = client.post(f"/api/documents/{document_id}/statement-analysis", json={})
        assert analysis.status_code == 200
        assert analysis.json()["status"] == "ready"
        assert analysis.json()["statement_account_id"]
        accounts = client.get("/api/statement-accounts")
        assert accounts.status_code == 200
        assert accounts.json()[0]["display_name"] == "中國信託信用卡"

        confirmation = client.post(
            f"/api/statements/{analysis.json()['statement_id']}/confirm",
            json={"review_version": analysis.json()["review_version"]},
        )
        assert confirmation.status_code == 200
        assert confirmation.json()["transaction_count"] == 1
        assert client.get("/api/finance/transactions").json()["total"] == 1


@pytest.mark.parametrize(("bank_id", "name"), [
    ("cathay", "國泰世華信用卡"), ("sinopac", "永豐信用卡"),
])
def test_new_bank_accounts_are_local_and_confirmation_remains_idempotent(tmp_path, bank_id, name):
    parser = FakeStatementParser()
    parser.bank_id = bank_id
    client, settings = make_client(tmp_path, parser)
    with client:
        document_id = upload_pdf(client)
        analyzed = client.post(f"/api/documents/{document_id}/statement-analysis", json={}).json()
        assert analyzed["status"] == "ready"
        assert client.get("/api/finance/transactions").json()["total"] == 0
        accounts = client.get("/api/statement-accounts").json()
        assert len(accounts) == 1
        assert accounts[0]["display_name"] == name
        repeated = client.post(f"/api/documents/{document_id}/statement-analysis", json={}).json()
        assert repeated["statement_id"] == analyzed["statement_id"]
        assert len(client.get("/api/statement-accounts").json()) == 1
        stale = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": analyzed["review_version"]})
        assert stale.status_code == 409
        confirmed = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": repeated["review_version"]})
        assert confirmed.status_code == 200
        again = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": repeated["review_version"]})
        assert again.json()["reused"]
        assert client.get("/api/finance/transactions").json()["total"] == 3
        assert client.post("/api/exports/excel/settings", json={"enabled": True}).json()["exported_transactions"] == 3
        with closing(load_workbook(settings.workbook_path, data_only=True)) as workbook:
            row = next(workbook["信用卡月支出"].iter_rows(min_row=2, values_only=True))
            assert row == (datetime(2026, 9, 1), name, "TWD", 120, 20, 100, 0, 0, 100, 300, 3)


def test_undated_interest_is_visible_but_cannot_be_confirmed(tmp_path):
    class UndatedInterestParser(FakeStatementParser):
        bank_id = "cathay"

        def parse(self, extracted_text, *, page_count):
            result = super().parse(extracted_text, page_count=page_count)
            undated = StatementLine(
                line_index=4, posting_date=date(2026, 9, 30), description="循環利息",
                transaction_kind="interest", amount=-5, currency="TWD",
            )
            return result.model_copy(update={"statement": result.statement.model_copy(update={"lines": (*result.statement.lines, undated)})})

    client, _settings = make_client(tmp_path, UndatedInterestParser())
    with client:
        document_id = upload_pdf(client)
        analyzed = client.post(f"/api/documents/{document_id}/statement-analysis", json={}).json()
        assert analyzed["status"] == "pending"
        assert analyzed["reason_code"] == "missing_transaction_date"
        assert analyzed["line_count"] == 4
        detail = client.get(f"/api/statements/{analyzed['statement_id']}").json()
        assert detail["lines"][-1]["transaction_date"] is None
        confirm = client.post(f"/api/statements/{analyzed['statement_id']}/confirm", json={"review_version": analyzed["review_version"]})
        assert confirm.status_code == 409
        assert client.get("/api/finance/transactions").json()["total"] == 0
