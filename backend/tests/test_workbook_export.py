from io import BytesIO
from threading import Event
import sys

from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, text
import pytest

from family_finance_hub.config import Settings
from family_finance_hub.database import make_engine
from family_finance_hub.exports.ports import WorkbookOwnershipError
from family_finance_hub.exports.xlsx import XlsxWorkbookWriter
from family_finance_hub.main import create_app
from family_finance_hub.models import ImportJob
from test_gmail_sync import FakeGmail, MemorySecretStore


def make_app(tmp_path, **kwargs):
    return create_app(Settings(
        database_url=f"sqlite:///{(tmp_path / 'export.db').as_posix()}",
        storage_root=tmp_path / "documents",
    ), create_schema=True, **kwargs)


def import_csv(client, description="Synthetic market"):
    csv = f"date,description,amount,currency\n2026-09-01,{description},-125.50,TWD\n2026-09-02,Refund,25.50,TWD\n2026-09-03,Book,-8,USD\n"
    response = client.post("/api/finance/import-csv", files={"file": ("synthetic.csv", csv.encode(), "text/csv")})
    assert response.status_code == 200
    return response.json()["document_id"]


def read_status(client):
    response = client.get("/api/exports/excel/status")
    assert response.status_code == 200
    return response.json()


def read_rows(client):
    response = client.get("/api/exports/excel/content")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        return {sheet.title: list(sheet.values) for sheet in workbook if sheet.sheet_state == "visible"}
    finally:
        workbook.close()


def test_export_disabled_by_default_then_rebuilds_active_transactions(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        status = read_status(client)
        assert status["enabled"] is False
        assert status["file_available"] is False
        assert app.state.workbook_export.refresh(scheduled=True) is False
        assert client.get("/api/exports/excel/content").status_code == 409
        document_id = import_csv(client)
        import_csv(client)
        status = client.post("/api/exports/excel/settings", json={"enabled": True}).json()
        assert status["status"] == "completed"
        assert status["exported_transactions"] == 3
        assert status["needs_update"] is False
        original = read_rows(client)
        assert "Synthetic market" in str(original)
        assert document_id in str(original)
        with app.state.session_factory() as session:
            jobs_before = len(session.scalars(select(ImportJob).where(ImportJob.source_type == "excel_export")).all())
        assert app.state.workbook_export.refresh(scheduled=True) is False
        with app.state.session_factory() as session:
            assert len(session.scalars(select(ImportJob).where(ImportJob.source_type == "excel_export")).all()) == jobs_before

        impact = client.get(f"/api/documents/{document_id}/import-impact").json()
        revoked = client.post(f"/api/documents/{document_id}/revoke", json={"impact_token": impact["impact_token"], "reason": "synthetic test"})
        assert revoked.status_code == 200
        assert read_status(client)["needs_update"] is True
        assert client.get("/api/exports/excel/content").status_code == 409
        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_status(client)["exported_transactions"] == 0
        assert "Synthetic market" not in str(read_rows(client))
        import_csv(client)
        assert app.state.workbook_export.refresh(scheduled=True) is False
        impact = client.get(f"/api/documents/{document_id}/import-impact").json()
        assert client.post(f"/api/documents/{document_id}/restore", json={"impact_token": impact["impact_token"]}).status_code == 200
        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_rows(client) == original


def test_manual_export_does_not_enable_automation_and_pdf_remains_pending(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        import_csv(client)
        assert client.post("/api/documents", files={"file": ("synthetic-locked.pdf", b"placeholder-not-parsed", "application/pdf")}).status_code == 200
        response = client.post("/api/exports/excel/refresh")
        assert response.status_code == 200
        status = response.json()
        assert not status["enabled"]
        assert status["pending_pdf_documents"] == 1
        assert not status["pdf_transaction_parser_ready"]
        assert status["exported_transactions"] == 3
        assert "synthetic-locked.pdf" in str(read_rows(client))


class FailingWriter:
    def __init__(self, error_type=PermissionError):
        self.error_type = error_type
        self.fail = False
        self.written = Event()

    def write(self, snapshot, destination):
        if self.fail:
            raise self.error_type("never expose synthetic secret or underlying path")
        XlsxWorkbookWriter().write(snapshot, destination)
        self.written.set()


def test_locked_workbook_retries_without_losing_import_or_old_file(tmp_path):
    writer = FailingWriter()
    app = make_app(tmp_path, workbook_writer=writer)
    with TestClient(app) as client:
        import_csv(client)
        client.post("/api/exports/excel/settings", json={"enabled": True})
        old_bytes = app.state.workbook_export.destination.read_bytes()
        old_success = read_status(client)["last_successful_at"]
        writer.fail = True
        import_csv(client, "Second statement")
        assert app.state.workbook_export.refresh(scheduled=True) is False
        status = read_status(client)
        assert status["status"] == "failed"
        assert status["last_error_code"] == "file_in_use"
        assert "synthetic secret" not in str(status)
        assert status["last_successful_at"] == old_success
        assert status["current_transactions"] == 6
        assert app.state.workbook_export.destination.read_bytes() == old_bytes
        assert app.state.workbook_export.refresh(scheduled=True) is False
        with app.state.session_factory() as session:
            failures = session.scalars(select(ImportJob).where(ImportJob.source_type == "excel_export", ImportJob.status == "failed")).all()
            assert len(failures) == 1
            assert "synthetic secret" not in failures[0].summary
        writer.fail = False
        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_status(client)["exported_transactions"] == 6
        assert read_status(client)["last_error"] is None


def test_restart_rebuilds_missing_workbook_and_disable_stops_updates(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        import_csv(client)
        client.post("/api/exports/excel/settings", json={"enabled": True})
        path = app.state.workbook_export.destination
    path.unlink()
    startup_writer = FailingWriter()
    second_app = make_app(tmp_path, workbook_writer=startup_writer)
    with TestClient(second_app) as client:
        assert startup_writer.written.wait(timeout=5)
        assert path.is_file()
        assert read_status(client)["enabled"]
        client.post("/api/exports/excel/settings", json={"enabled": False})
        import_csv(client, "While paused")
        assert second_app.state.workbook_export.refresh(scheduled=True) is False
        assert read_status(client)["needs_update"]
        assert read_status(client)["exported_transactions"] == 3


def test_gmail_to_workbook_is_idempotent_and_pdf_not_silently_counted(tmp_path):
    gmail = FakeGmail()
    app = make_app(tmp_path, secret_store=MemorySecretStore(), gmail_client_factory=lambda: gmail)
    with TestClient(app) as client:
        client.post("/api/exports/excel/settings", json={"enabled": True})
        synced = client.post("/api/gmail/sync", json={})
        assert synced.status_code == 200
        assert app.state.workbook_export.refresh(scheduled=True)
        status = read_status(client)
        assert status["exported_transactions"] == 1
        assert status["pending_pdf_documents"] == 1
        first = read_rows(client)
        assert client.post("/api/gmail/sync", json={}).json()["created_transactions"] == 0
        assert app.state.workbook_export.refresh(scheduled=True) is False
        assert read_rows(client) == first


def test_export_failure_does_not_advance_fingerprint(tmp_path):
    writer = FailingWriter(WorkbookOwnershipError)
    writer.fail = True
    app = make_app(tmp_path, workbook_writer=writer)
    with TestClient(app) as client:
        status = client.post("/api/exports/excel/settings", json={"enabled": True}).json()
        assert status["last_error_code"] == "unowned_workbook"
        assert status["last_successful_at"] is None
        writer.fail = False
        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_status(client)["needs_update"] is False


def test_refresh_lock_reports_conflict(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client, app.state.workbook_export.lock:
        assert app.state.workbook_export.refresh(scheduled=True) is False
        assert client.post("/api/exports/excel/refresh").status_code == 409


def test_workbook_migration_defaults_off_and_preserves_documents(tmp_path):
    path = tmp_path / "migration.db"
    config = Config("backend/alembic.ini")
    config.set_main_option("sqlalchemy.url", f"sqlite:///{path.as_posix()}")
    command.upgrade(config, "0009_document_revocation")
    engine = make_engine(f"sqlite:///{path.as_posix()}")
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO documents (id, sha256, filename, content_type, size_bytes, created_at) VALUES ('doc', 'hash', 'synthetic.pdf', 'application/pdf', 10, '2026-09-01')"))
    command.upgrade(config, "head")
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO workbook_export_state (id) VALUES ('excel')"))
        assert connection.execute(text("SELECT enabled, status, transaction_count, output_sha256 FROM workbook_export_state")).one() == (0, "never", 0, None)
        assert connection.scalar(text("SELECT filename FROM documents WHERE id = 'doc'")) == "synthetic.pdf"
    command.downgrade(config, "0009_document_revocation")
    with engine.connect() as connection:
        assert connection.scalar(text("SELECT COUNT(*) FROM documents")) == 1
    engine.dispose()


def test_destination_change_requires_new_export(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        import_csv(client)
        client.post("/api/exports/excel/refresh")
        service = app.state.workbook_export
        service.destination = tmp_path / "different.xlsx"
        assert read_status(client)["needs_update"]
        assert client.get("/api/exports/excel/content").status_code == 409
        client.post("/api/exports/excel/refresh")
        assert read_status(client)["needs_update"] is False


def test_external_workbook_edit_is_not_served_and_is_rebuilt(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        import_csv(client)
        assert client.post("/api/exports/excel/settings", json={"enabled": True}).status_code == 200
        path = app.state.workbook_export.destination
        workbook = load_workbook(path)
        workbook["月份幣別摘要"]["A1"] = "手動修改"
        workbook.save(path)
        workbook.close()

        status = read_status(client)
        assert status["needs_update"] is True
        assert status["file_available"] is False
        assert client.get("/api/exports/excel/content").status_code == 409

        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_status(client)["needs_update"] is False
        assert read_rows(client)["月份幣別摘要"][0][0] == "月份"


def test_export_includes_all_rows_beyond_ui_page_and_current_month(tmp_path):
    app = make_app(tmp_path)
    with TestClient(app) as client:
        rows = [f"2025-08-{index % 28 + 1:02},Synthetic row {index},-1.25,TWD" for index in range(260)]
        content = ("date,description,amount,currency\n" + "\n".join(rows)).encode()
        assert client.post("/api/finance/import-csv", files={"file": ("large.csv", content, "text/csv")}).status_code == 200
        assert client.post("/api/exports/excel/refresh").json()["exported_transactions"] == 260
        assert len(read_rows(client)["交易明細"]) == 261
        assert read_rows(client)["月份幣別摘要"][1][3] == 325


@pytest.mark.skipif(sys.platform != "win32", reason="Windows sharing semantics")
def test_windows_file_lock_preserves_workbook_until_released(tmp_path):
    import ctypes
    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    kernel32.CreateFileW.restype = wintypes.HANDLE
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    app = make_app(tmp_path)
    with TestClient(app) as client:
        import_csv(client)
        client.post("/api/exports/excel/settings", json={"enabled": True})
        path = app.state.workbook_export.destination
        old_bytes = path.read_bytes()
        handle = kernel32.CreateFileW(str(path), 0x80000000, 1, None, 3, 0, None)
        assert handle != wintypes.HANDLE(-1).value
        try:
            import_csv(client, "While workbook is open")
            assert app.state.workbook_export.refresh(scheduled=True) is False
            assert read_status(client)["last_error_code"] == "file_in_use"
            assert path.read_bytes() == old_bytes
            assert list(path.parent.glob("*.tmp")) == []
        finally:
            kernel32.CloseHandle(handle)
        assert app.state.workbook_export.refresh(scheduled=True)
        assert read_status(client)["exported_transactions"] == 6
