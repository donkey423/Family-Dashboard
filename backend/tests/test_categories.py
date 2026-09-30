from datetime import date
from contextlib import closing
from decimal import Decimal
from io import BytesIO
import random
import shutil
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from time import perf_counter

from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, event, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from family_finance_hub.config import Settings
from family_finance_hub.exports.service import read_snapshot, WorkbookExportService
from family_finance_hub.exports.xlsx import XlsxWorkbookWriter
from family_finance_hub.finance.categories.service import CategorizationService, seed_categories
from family_finance_hub.finance.queries import transaction_totals
from family_finance_hub.main import create_app
from family_finance_hub.models import Document, FinanceCategory, FinanceCategoryRule, FinanceTransaction, Statement, TransactionCategoryOverride, utc_now


@pytest.fixture
def client(tmp_path):
    app = create_app(Settings(database_url=f"sqlite:///{(tmp_path / 'category.db').as_posix()}", storage_root=tmp_path / "documents"), create_schema=True)
    with TestClient(app) as value:
        with app.state.session_factory() as session, session.begin():
            session.add(Document(id="doc", sha256="c" * 64, filename="synthetic.csv", content_type="text/csv", size_bytes=0))
            session.flush()
            session.add(Statement(id="statement", document_id="doc", bank_id="synthetic", format_version="1", status="imported"))
            session.flush()
            data = [("a", "  Shop 123 ", "-100", None, None), ("b", "ＳＨＯＰ 123", "-50", None, None),
                    ("c", "Coffee", "-30", "statement", "purchase"), ("d", "Coffee", "10", "statement", "refund"),
                    ("e", "Interest", "-5", "statement", "interest"), ("f", "Payment", "1000", "statement", "payment"),
                    ("g", "Salary", "5000", None, None), ("h", "Refund Only", "40", "statement", "refund")]
            for index, (key, name, amount, statement, kind) in enumerate(data):
                session.add(FinanceTransaction(id=key, source_document_id="doc", row_hash=key, transaction_date=date(2026, 9, index + 1), description=name, amount=Decimal(amount), currency="TWD", raw_json="{}", statement_id=statement, transaction_kind=kind, statement_line_index=index if statement else None))
        yield value


def get_spending(client, **params):
    response = client.get("/api/finance/spending-by-category", params={"month": "2026-09", "currency": "TWD", **params})
    assert response.status_code == 200
    return response.json()


def assign(client, key, category, scope="transaction"):
    return client.patch(f"/api/finance/transactions/{key}/category", json={"category_id": category, "scope": scope})


def test_resolution_updates_filters_merchants_and_preserves_identity(client):
    before = client.get("/api/dashboard", params={"month": "2026-09"}).json()
    with client.app.state.session_factory() as session:
        identity = [(row.id, row.row_hash, row.statement_id, row.statement_line_index, row.raw_json) for row in session.scalars(select(FinanceTransaction).order_by(FinanceTransaction.id))]
    assert assign(client, "b", "travel").status_code == 200
    assert assign(client, "a", "shopping", "merchant").status_code == 200
    rows = client.get("/api/finance/transactions", params={"category_id": "shopping", "limit": 1}).json()
    assert rows["total"] == 1 and rows["items"][0]["id"] == "a"
    assert assign(client, "b", "food", "merchant").status_code == 200
    assert client.get("/api/finance/transactions", params={"category_id": "travel"}).json()["total"] == 1
    assert client.delete("/api/finance/transactions/b/category").json()["category_source"] == "exact_rule"
    assert client.get("/api/finance/transactions", params={"category_id": "food", "limit": 1, "offset": 1}).json()["total"] == 2
    merchants = client.get("/api/finance/category-merchants", params={"category_id": "food"}).json()["items"]
    assert merchants[0]["net_amount"] == "150.00" and merchants[0]["transaction_count"] == 2
    assert client.get("/api/finance/transactions", params={"merchant_key": "shop 123"}).json()["total"] == 2
    after = client.get("/api/dashboard", params={"month": "2026-09"}).json()
    assert before == after
    with client.app.state.session_factory() as session:
        assert identity == [(row.id, row.row_hash, row.statement_id, row.statement_line_index, row.raw_json) for row in session.scalars(select(FinanceTransaction).order_by(FinanceTransaction.id))]


def test_rule_crud_disabled_inactive_fallback_and_duplicate(client):
    response = client.post("/api/finance/category-rules", json={"category_id": "food", "match_type": "contains", "pattern": "shop"})
    assert response.status_code == 201
    rule = response.json()
    assert client.post("/api/finance/category-rules", json={"category_id": "travel", "match_type": "contains", "pattern": "ＳＨＯＰ"}).status_code == 409
    assert client.patch(f"/api/finance/category-rules/{rule['id']}", json={"enabled": False}).status_code == 200
    assert client.get("/api/finance/transactions", params={"category_id": "food"}).json()["total"] == 0
    assert client.patch(f"/api/finance/category-rules/{rule['id']}", json={"enabled": True}).status_code == 200
    assert client.patch("/api/finance/categories/food", json={"is_active": False}).status_code == 200
    assert assign(client, "a", "food").status_code == 422
    assert client.get("/api/finance/transactions", params={"category_id": "food"}).json()["total"] == 0
    assert client.delete(f"/api/finance/category-rules/{rule['id']}").status_code == 200
    assert client.delete(f"/api/finance/category-rules/{rule['id']}").status_code == 404


@pytest.mark.parametrize("path,params", [
    ("spending-by-category", {"month": "2026-13"}), ("spending-by-category", {"currency": "12!"}),
    ("transactions", {"category_id": "missing"}), ("category-merchants", {"category_id": "missing"}),
])
def test_invalid_filters_fail_closed(client, path, params):
    response = client.get(f"/api/finance/{path}", params=params)
    assert response.status_code in (404, 422)
    assert "raw_json" not in response.text


@pytest.mark.parametrize("code", ["uncategorized", "income", "transfer"])
def test_required_categories_protected(client, code):
    assert client.patch(f"/api/finance/categories/{code}", json={"is_active": False}).status_code == 422
    assert client.post("/api/finance/categories", json={"code": code, "display_name": "Duplicate"}).status_code == 409


@pytest.mark.parametrize("transaction,category", [("a", "missing"), ("missing", "food")])
def test_runtime_rejects_invalid_override_foreign_keys(client, transaction, category):
    with client.app.state.session_factory() as session:
        assert session.scalar(text("PRAGMA foreign_keys")) == 1
        session.add(TransactionCategoryOverride(transaction_id=transaction, category_id=category))
        with pytest.raises(Exception, match="FOREIGN KEY"):
            session.flush()
        session.rollback()


def test_category_validation_and_rename(client):
    assert client.post("/api/finance/categories", json={"code": "new", "display_name": " "}).status_code == 422
    created = client.post("/api/finance/categories", json={"code": "new", "display_name": "New"}).json()
    assert assign(client, "a", created["id"]).status_code == 200
    assert client.patch(f"/api/finance/categories/{created['id']}", json={"display_name": "Renamed"}).status_code == 200
    assert "Renamed" in str(get_spending(client))
    assert assign(client, "missing", "food").status_code == 404
    assert client.post("/api/finance/category-rules", json={"category_id": "food", "match_type": "contains", "pattern": " "}).status_code == 422
    assert client.post("/api/finance/categories", json={"code": "new", "display_name": "Other"}).status_code == 409
    assert client.post("/api/finance/categories", json={"code": "another", "display_name": "Renamed"}).status_code == 201


def test_inactive_category_preserves_references_and_restores_override(client):
    assert assign(client, "a", "food", "merchant").status_code == 200
    assert assign(client, "b", "food").status_code == 200
    assert client.patch("/api/finance/categories/food", json={"is_active": False}).status_code == 200
    assert client.get("/api/finance/transactions", params={"category_id": "food"}).json()["total"] == 0
    with client.app.state.session_factory() as session:
        assert session.get(TransactionCategoryOverride, "b").category_id == "food"
        assert session.scalar(select(FinanceCategoryRule.category_id).where(FinanceCategoryRule.normalized_pattern == "SHOP 123")) == "food"
        assert not session.execute(text("PRAGMA foreign_key_check")).all()
    assert client.patch("/api/finance/categories/food", json={"is_active": True}).status_code == 200
    rows = client.get("/api/finance/transactions", params={"category_id": "food"}).json()["items"]
    assert {row["id"]: row["category_source"] for row in rows} == {"a": "exact_rule", "b": "override"}


def test_refunds_payment_income_excel_parity_and_fingerprint(client, tmp_path):
    assert assign(client, "h", "travel").status_code == 200
    assert assign(client, "f", "food").status_code == 200  # Label never changes payment semantics.
    summary = get_spending(client)
    assert summary["dashboard_expense"] == "135.00"
    assert summary["positive_category_total"] == "175.00"
    assert summary["refund_credit_total"] == "-40.00"
    assert sum(Decimal(row["net_amount"]) for row in summary["categories"]) == Decimal("135")
    with client.app.state.session_factory() as session:
        snapshot = read_snapshot(session)
    destination = tmp_path / "categories.xlsx"
    XlsxWorkbookWriter().write(snapshot, destination)
    with closing(load_workbook(destination)) as workbook:
        values = list(workbook["分類支出"].values)[1:]
        assert sum(Decimal(str(row[3])) for row in values) == Decimal("135")
        assert sum(row[4] for row in values) == 6
        assert workbook["交易明細"].cell(1, 16).value == "分類"
    assert client.post("/api/finance/category-rules", json={"category_id": "food", "match_type": "contains", "pattern": "unused"}).status_code == 201
    with client.app.state.session_factory() as session:
        changed = read_snapshot(session)
    assert changed.transactions == snapshot.transactions
    assert changed.category_configuration_hash != snapshot.category_configuration_hash


def test_lifecycle_currencies_dates_restore_reclassifies(client):
    with client.app.state.session_factory() as session, session.begin():
        session.get(FinanceTransaction, "a").currency = "USD"
        session.get(FinanceTransaction, "b").transaction_date = None
        session.get(FinanceTransaction, "c").transaction_date = date(2024, 2, 29)
    assert get_spending(client, currency="USD")["dashboard_expense"] == "100.00"
    assert get_spending(client, month="2024-02")["dashboard_expense"] == "30.00"
    assert get_spending(client)["dashboard_expense"] == "-45.00"
    assert assign(client, "b", "shopping").status_code == 200
    impact = client.get("/api/documents/doc/import-impact").json()
    assert client.post("/api/documents/doc/revoke", json={"impact_token": impact["impact_token"]}).status_code == 200
    assert get_spending(client)["dashboard_expense"] == "0.00"
    assert assign(client, "a", "food").status_code == 404
    assert client.post("/api/finance/category-rules", json={"category_id": "food", "match_type": "normalized_exact", "pattern": "Shop 123"}).status_code == 201
    impact = client.get("/api/documents/doc/import-impact").json()
    assert client.post("/api/documents/doc/restore", json={"impact_token": impact["impact_token"]}).status_code == 200
    assert client.get("/api/finance/transactions", params={"category_id": "shopping"}).json()["total"] == 1
    assert get_spending(client, currency="USD")["categories"][0]["code"] == "food"


def test_same_effective_override_changes_excel_fingerprint(client):
    with client.app.state.session_factory() as session:
        before = read_snapshot(session)
    assert assign(client, "a", "uncategorized").status_code == 200
    with client.app.state.session_factory() as session:
        after = read_snapshot(session)
    assert after.transactions == before.transactions
    assert after.category_configuration_hash != before.category_configuration_hash
    assert client.delete("/api/finance/transactions/a/category").status_code == 200
    with client.app.state.session_factory() as session:
        cleared = read_snapshot(session)
    assert cleared.category_configuration_hash == before.category_configuration_hash


def test_merchant_rule_applies_to_future_import_without_overwriting_exception(client):
    assert assign(client, "a", "food", "merchant").status_code == 200
    assert assign(client, "b", "travel").status_code == 200
    file = {"file": ("future.csv", b"date,description,amount,currency\n2026-09-25,Shop 123,-25,TWD\n", "text/csv")}
    imported = client.post("/api/finance/import-csv", files=file)
    assert imported.status_code == 200
    assert client.post("/api/finance/import-csv", files=file).json()["created_transactions"] == 0
    food = client.get("/api/finance/transactions", params={"category_id": "food"}).json()
    assert food["total"] == 2
    assert client.get("/api/finance/transactions", params={"category_id": "travel"}).json()["total"] == 1


def test_twenty_merchant_rows_keep_three_manual_exceptions(client):
    with client.app.state.session_factory() as session, session.begin():
        for index in range(20):
            session.add(FinanceTransaction(id=f"bulk{index}", source_document_id="doc", row_hash=f"bulk{index}", description="Bulk shop", amount=Decimal("-1"), currency="TWD", raw_json="{}", transaction_date=date(2026, 9, 30)))
    for index in range(3):
        assert assign(client, f"bulk{index}", "travel").status_code == 200
    assert assign(client, "bulk3", "food", "merchant").status_code == 200
    rows = client.get("/api/finance/transactions", params={"merchant_key": "Bulk shop"}).json()["items"]
    assert sum(row["category_code"] == "food" for row in rows) == 17
    assert sum(row["category_code"] == "travel" for row in rows) == 3
    assert client.get("/api/finance/transactions", params={"merchant_key": "missing"}).json()["total"] == 0


@pytest.mark.parametrize("scope", ["transaction", "merchant"])
def test_failed_write_rolls_back_previous_category(client, monkeypatch, scope):
    assert assign(client, "a", "food", scope).status_code == 200
    before = get_spending(client)
    with monkeypatch.context() as patch:
        def reject_commit(_session):
            raise IntegrityError("synthetic", {}, Exception("synthetic conflict"))
        patch.setattr(Session, "commit", reject_commit)
        assert assign(client, "a", "travel", scope).status_code == 409
    assert get_spending(client) == before
    assert client.get("/api/finance/transactions", params={"category_id": "travel"}).json()["total"] == 0


def test_locked_database_returns_retryable_conflict_without_partial_write(client):
    factory = client.app.state.session_factory
    engine = factory.kw["bind"]
    def short_timeout(connection, *_args):
        connection.execute("PRAGMA busy_timeout=1")
    event.listen(engine, "checkout", short_timeout)
    try:
        with closing(sqlite3.connect(engine.url.database)) as blocker:
            blocker.execute("BEGIN IMMEDIATE")
            result = assign(client, "a", "travel")
            assert result.status_code == 409
            assert "sql" not in result.text.lower()
            blocker.rollback()
        assert client.get("/api/finance/transactions", params={"category_id": "travel"}).json()["total"] == 0
        assert assign(client, "a", "food").status_code == 200
    finally:
        event.remove(engine, "checkout", short_timeout)


def test_concurrent_writes_have_one_effective_override(client):
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda category: (category, assign(client, "a", category).status_code), ["food", "travel"]))
    assert all(status in (200, 409) for _, status in results)
    successful = {category for category, status in results if status == 200}
    assert successful
    with client.app.state.session_factory() as session:
        overrides = session.scalars(select(TransactionCategoryOverride).where(TransactionCategoryOverride.transaction_id == "a")).all()
        assert len(overrides) == 1 and overrides[0].category_id in successful
    assert get_spending(client)["dashboard_expense"] == "135.00"


@pytest.mark.parametrize("purchase,refunds,expected", [
    ("100", ["20"], "80.00"), ("100", ["100"], "0.00"),
    ("100", ["150"], "-50.00"), ("0", ["40"], "-40.00"),
    ("100", ["20", "30"], "50.00"),
])
def test_refund_edges_preserve_signed_net(client, purchase, refunds, expected):
    with client.app.state.session_factory() as session, session.begin():
        session.get(FinanceTransaction, "c").amount = -Decimal(purchase)
        session.get(FinanceTransaction, "d").amount = Decimal(refunds[0])
        for index, value in enumerate(refunds[1:]):
            session.add(FinanceTransaction(id=f"refund{index}", source_document_id="doc", row_hash=f"refund{index}", description="Coffee", amount=Decimal(value), currency="TWD", raw_json="{}", transaction_date=date(2026, 9, 30), statement_id="statement", transaction_kind="refund", statement_line_index=20 + index))
    assert assign(client, "c", "food", "merchant").status_code == 200
    summary = get_spending(client)
    food = next(row for row in summary["categories"] if row["code"] == "food")
    assert food["net_amount"] == expected
    merchants = client.get("/api/finance/category-merchants", params={"category_id": "food"}).json()["items"]
    assert sum(Decimal(row["net_amount"]) for row in merchants) == Decimal(expected)
    assert sum(Decimal(row["net_amount"]) for row in summary["categories"]) == Decimal(summary["dashboard_expense"])


def test_date_boundaries_and_three_currencies_do_not_mix(client):
    with client.app.state.session_factory() as session, session.begin():
        session.get(FinanceTransaction, "a").transaction_date = date(2026, 9, 1)
        session.get(FinanceTransaction, "b").transaction_date = date(2026, 9, 30)
        session.get(FinanceTransaction, "c").transaction_date = date(2026, 10, 1)
        session.get(FinanceTransaction, "e").transaction_date = date(2026, 8, 31)
        session.get(FinanceTransaction, "d").currency = "USD"
        session.get(FinanceTransaction, "h").currency = "JPY"
    expected = {"TWD": "150.00", "USD": "-10.00", "JPY": "-40.00"}
    for currency, amount in expected.items():
        assert get_spending(client, currency=currency)["dashboard_expense"] == amount
    assert get_spending(client, month="2026-10")["dashboard_expense"] == "30.00"
    assert get_spending(client, month="2026-08")["dashboard_expense"] == "5.00"


def test_excel_dirty_formula_name_currency_lifecycle_and_restore(client, tmp_path):
    service = WorkbookExportService(client.app.state.session_factory, XlsxWorkbookWriter(), tmp_path / "projection.xlsx")
    assert service.refresh()
    assert service.status()["needs_update"] is False
    assert assign(client, "a", "food").status_code == 200
    assert service.status()["needs_update"] is True
    assert client.patch("/api/finance/categories/food", json={"display_name": "=1+1"}).status_code == 200
    with client.app.state.session_factory() as session, session.begin():
        session.get(FinanceTransaction, "b").currency = "USD"
        session.get(FinanceTransaction, "h").currency = "JPY"
    assert service.refresh()
    baseline = service.read_workbook()
    with closing(load_workbook(BytesIO(baseline))) as workbook:
        cell = next(row[15] for row in workbook["交易明細"].iter_rows(min_row=2) if row[5].value == "a")
        assert cell.value == "=1+1" and cell.data_type == "s"
        totals = list(workbook["分類支出"].values)[1:]
        assert {row[1] for row in totals} == {"TWD", "USD", "JPY"}
        for currency in ("TWD", "USD", "JPY"):
            assert sum(Decimal(str(row[3])) for row in totals if row[1] == currency) == Decimal(get_spending(client, currency=currency)["dashboard_expense"])
    impact = client.get("/api/documents/doc/import-impact").json()
    client.post("/api/documents/doc/revoke", json={"impact_token": impact["impact_token"]})
    assert service.status()["needs_update"] is True and service.refresh()
    with closing(load_workbook(BytesIO(service.read_workbook()))) as workbook:
        assert workbook["分類支出"].max_row == 1
    impact = client.get("/api/documents/doc/import-impact").json()
    client.post("/api/documents/doc/restore", json={"impact_token": impact["impact_token"]})
    assert service.refresh()
    with closing(load_workbook(BytesIO(service.read_workbook()))) as workbook:
        assert list(workbook["分類支出"].values)[1:] == totals


def test_migration_forward_downgrade_preserves_rows_and_constraints(tmp_path):
    path = tmp_path / "migration.db"
    config = Config("backend/alembic.ini")
    config.set_main_option("sqlalchemy.url", f"sqlite:///{path.as_posix()}")
    command.upgrade(config, "0011_statement_import")
    engine = create_engine(f"sqlite:///{path.as_posix()}")
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO documents(id,sha256,filename,content_type,size_bytes,created_at) VALUES('doc','sha','file','text/csv',0,CURRENT_TIMESTAMP)"))
        connection.execute(text("INSERT INTO finance_transactions(id,source_document_id,row_hash,description,amount,currency,raw_json,created_at) VALUES('t','doc','hash','Shop',-1,'TWD','{}',CURRENT_TIMESTAMP)"))
        connection.execute(text("INSERT INTO statements(id,document_id,bank_id,format_version,status,review_version,created_at,updated_at) VALUES('s','doc','synthetic','v1','imported',1,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"))
        connection.execute(text("INSERT INTO finance_transactions(id,source_document_id,row_hash,description,amount,currency,raw_json,created_at,statement_id,statement_line_index,transaction_kind) VALUES('card','doc','cardhash','Shop',-2,'TWD','{}',CURRENT_TIMESTAMP,'s',1,'purchase')"))
        original = connection.execute(text("SELECT * FROM finance_transactions ORDER BY id")).all()
    command.upgrade(config, "head")
    command.upgrade(config, "head")
    with engine.connect() as connection:
        assert connection.scalar(text("SELECT count(*) FROM finance_categories")) == 14
        assert connection.execute(text("SELECT * FROM finance_transactions ORDER BY id")).all() == original
        connection.exec_driver_sql("PRAGMA foreign_keys=ON")
        with pytest.raises(Exception, match="FOREIGN KEY"):
            connection.execute(text("INSERT INTO transaction_category_overrides VALUES('t','missing',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"))
        connection.rollback()
        with pytest.raises(Exception, match="CHECK"):
            connection.execute(text("INSERT INTO finance_category_rules VALUES('bad','food','invalid','X',0,1,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"))
        connection.rollback()
        with pytest.raises(Exception, match="FOREIGN KEY"):
            connection.execute(text("INSERT INTO finance_category_rules VALUES('bad','missing','contains','X',0,1,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"))
    command.downgrade(config, "0011_statement_import")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT * FROM finance_transactions ORDER BY id")).all() == original
    command.upgrade(config, "head")
    engine.dispose()


def test_seed_idempotent_backup_restore_and_no_reseed(client, tmp_path):
    assign(client, "a", "food", "merchant")
    assign(client, "b", "travel")
    client.patch("/api/finance/categories/food", json={"display_name": "Custom food"})
    before = get_spending(client)
    with client.app.state.session_factory() as session:
        original_snapshot = read_snapshot(session)
    with client.app.state.session_factory() as session, session.begin():
        seed_categories(session)
        seed_categories(session)
        assert len(session.scalars(select(FinanceCategory)).all()) == 14
    source = tmp_path / "category.db"
    restored = tmp_path / "restore.db"
    shutil.copy2(source, restored)
    app = create_app(Settings(database_url=f"sqlite:///{restored.as_posix()}", storage_root=tmp_path / "restore-documents"))
    with TestClient(app) as other:
        assert get_spending(other) == before
        assert len(other.get("/api/finance/category-rules").json()) == 1
        assert other.get("/api/finance/transactions", params={"category_id": "travel"}).json()["total"] == 1
        with other.app.state.session_factory() as session:
            restored_snapshot = read_snapshot(session)
        assert restored_snapshot.transactions == original_snapshot.transactions
        assert restored_snapshot.category_configuration_hash == original_snapshot.category_configuration_hash
        for name, snapshot in (("before", original_snapshot), ("restored", restored_snapshot)):
            XlsxWorkbookWriter().write(snapshot, tmp_path / f"{name}.xlsx")
        with closing(load_workbook(tmp_path / "before.xlsx")) as original, closing(load_workbook(tmp_path / "restored.xlsx")) as recovered:
            for sheet in ("交易明細", "分類支出"):
                assert list(original[sheet].values) == list(recovered[sheet].values)


def test_batch_query_count_randomized_invariants_and_10000_rows(client, capsys):
    rng = random.Random(423)
    with client.app.state.session_factory() as session, session.begin():
        for index in range(10000):
            session.add(FinanceTransaction(id=f"r{index}", source_document_id="doc", row_hash=f"r{index}", description=f"SHOP {index % 50}", amount=Decimal(-rng.randrange(1, 1000)) / 100, currency="TWD", raw_json="{}", transaction_date=date(2026, 9, 1)))
        for index in range(50):
            session.add(FinanceCategoryRule(id=f"rule{index}", category_id="food", match_type="normalized_exact", normalized_pattern=f"SHOP {index}", priority=0, enabled=True))
    with client.app.state.session_factory() as session:
        rows = session.scalars(select(FinanceTransaction)).all()
        queries = []
        def observe(*args):
            queries.append(args[2])
        engine = session.get_bind()
        event.listen(engine, "before_cursor_execute", observe)
        started = perf_counter()
        service = CategorizationService(session)
        resolved = service.resolve(rows)
        spending = service.spending(rows, resolved)
        elapsed = (perf_counter() - started) * 1000
        event.remove(engine, "before_cursor_execute", observe)
        assert len(queries) == 3
        assert spending == service.spending(rows, service.resolve(rows))
        assert spending["dashboard_expense"] == transaction_totals(session)["currency_totals"][0]["expenses"]
        print(f"CATEGORY_BENCHMARK 10008 rows 50 rules {elapsed:.1f}ms, 3 configuration/override queries")
    api_times = []
    for _ in range(3):
        started = perf_counter()
        response = client.get("/api/finance/spending-by-category", params={"month": "2026-09", "currency": "TWD"})
        api_times.append((perf_counter() - started) * 1000)
        assert response.status_code == 200
        assert response.json() == {"month": "2026-09", "currency": "TWD", **spending}
    print("CATEGORY_API_BENCHMARK 10008 rows 50 rules " + ", ".join(f"{value:.1f}ms" for value in api_times))


def test_500_merchant_api_and_category_pagination_use_batch_queries(client):
    with client.app.state.session_factory() as session, session.begin():
        session.add(FinanceCategoryRule(id="bulk", category_id="food", match_type="contains", normalized_pattern="BULK", priority=0, enabled=True))
        for index in range(1000):
            session.add(FinanceTransaction(id=f"merchant{index:04}", source_document_id="doc", row_hash=f"merchant{index:04}",
                description=f"Bulk {index % 500:03}", amount=Decimal("-1"), currency="TWD", raw_json="{\"private_source\":true}", transaction_date=date(2026, 9, 30)))
    engine = client.app.state.session_factory.kw["bind"]
    queries = []
    def observe(*args):
        queries.append(args[2])
    event.listen(engine, "before_cursor_execute", observe)
    try:
        response = client.get("/api/finance/category-merchants", params={"category_id": "food", "month": "2026-09", "currency": "TWD"})
        assert response.status_code == 200
        merchants = response.json()["items"]
        assert len(merchants) == 500 and all(row["transaction_count"] == 2 for row in merchants)
        assert sum(Decimal(row["net_amount"]) for row in merchants) == Decimal("1000")
        assert "private_source" not in response.text and "raw_json" not in response.text
        first_name = merchants[0]["display_name"]
        assert len(queries) <= 6
        queries.clear()
        response = client.get("/api/finance/transactions", params={"category_id": "food", "month": "2026-09", "currency": "TWD", "limit": 50, "offset": 950})
        assert response.status_code == 200
        assert response.json()["total"] == 1000 and len(response.json()["items"]) == 50
        assert len(queries) <= 6
        queries.clear()
        repeated = client.get("/api/finance/category-merchants", params={"category_id": "food", "month": "2026-09", "currency": "TWD"})
        assert repeated.json()["items"][0]["display_name"] == first_name
    finally:
        event.remove(engine, "before_cursor_execute", observe)
