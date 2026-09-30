import os
import traceback
from contextlib import closing
from dataclasses import FrozenInstanceError, replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from family_finance_hub.exports import (
    ExportDocument,
    ExportTransaction,
    WorkbookOwnershipError,
    WorkbookSnapshot,
    WorkbookWriteError,
    WorkbookWriter,
    XlsxWorkbookWriter,
)
from family_finance_hub.exports import xlsx


def transaction(
    id="000001", *, amount="-12.50", currency="TWD", transaction_date=date(2026, 9, 23)
):
    return ExportTransaction(
        id=id,
        document_id="000042",
        transaction_date=transaction_date,
        description="合成測試交易",
        amount=Decimal(amount),
        currency=currency,
        source_filename="九月帳單.csv",
    )


def empty_snapshot():
    return WorkbookSnapshot(transactions=(), documents=())


def test_snapshot_contract_is_frozen_and_preserves_stable_ids(tmp_path):
    tx = transaction()
    document = ExportDocument("000042", "九月帳單.csv", "active", 1)
    snapshot = WorkbookSnapshot((tx,), (document,))
    for instance, field, value in (
        (tx, "id", "changed"),
        (document, "id", "changed"),
        (snapshot, "transactions", ()),
    ):
        with pytest.raises(FrozenInstanceError):
            setattr(instance, field, value)

    destination = tmp_path / "家庭收支記錄.xlsx"
    writer: WorkbookWriter = XlsxWorkbookWriter()
    assert writer.write(snapshot, destination) is None
    writer.write(snapshot, destination)
    with closing(load_workbook(destination)) as workbook:
        detail = workbook["交易明細"]
        assert tuple(cell.value for cell in detail[1]) == (
            "交易日期", "說明", "金額", "幣別", "來源檔名", "交易 ID", "文件 ID",
            "來源類型", "帳戶", "交易類型", "入帳日期", "帳單起日", "帳單迄日", "帳單列號", "Statement ID",
            "分類",
        )
        assert detail["F2"].value == tx.id
        assert detail["G2"].value == document.id
        assert detail["F2"].data_type == detail["G2"].data_type == "s"
        assert workbook["文件狀態"]["A2"].value == document.id
    assert list(tmp_path.iterdir()) == [destination]


def test_summaries_separate_months_and_currencies_with_numeric_dates_and_amounts(tmp_path):
    snapshot = WorkbookSnapshot((
        transaction("jpy", amount="-200", currency="JPY"),
        transaction("oct", amount="-2.05", transaction_date=date(2026, 10, 1)),
        transaction("usd-in", amount="50.20", currency="USD"),
        transaction("usd-out", amount="-3.10", currency="USD"),
        transaction("twd-in", amount="1000.10"),
        transaction("twd-out", amount="-100.05"),
        transaction("zero", amount="0"),
    ), ())
    destination = tmp_path / "summary.xlsx"
    XlsxWorkbookWriter().write(snapshot, destination)
    with closing(load_workbook(destination, data_only=False)) as workbook:
        summary = workbook["月份幣別摘要"]
        assert list(summary.iter_rows(min_row=2, values_only=True)) == [
            (datetime(2026, 9, 1), "JPY", 0, 200, -200, 1),
            (datetime(2026, 9, 1), "TWD", 1000.10, 100.05, 900.05, 3),
            (datetime(2026, 9, 1), "USD", 50.20, 3.10, 47.10, 2),
            (datetime(2026, 10, 1), "TWD", 0, 2.05, -2.05, 1),
        ]
        assert summary["A2"].number_format == "yyyy-mm"
        for row in summary.iter_rows(min_row=2, min_col=3):
            assert all(cell.data_type == "n" for cell in row)
        detail = workbook["交易明細"]
        assert detail.max_row == len(snapshot.transactions) + 1
        assert detail["A2"].value == datetime(2026, 9, 23)
        assert detail["A2"].number_format == "yyyy-mm-dd"
        assert all(row[0].data_type == "n" for row in detail.iter_rows(min_row=2, min_col=3, max_col=3))

    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(destination) as archive:
        sheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        cell = sheet.find(".//x:c[@r='A2']", namespace)
        assert cell.attrib["t"] == "n"
        assert float(cell.find("x:v", namespace).text) > 0


def test_missing_dates_are_preserved_and_grouped_by_currency_in_unknown_month(tmp_path):
    snapshot = WorkbookSnapshot((
        transaction("known", amount="5"),
        transaction("unknown-twd-in", transaction_date=None, amount="10.25"),
        transaction("unknown-twd-out", transaction_date=None, amount="-0.25"),
        transaction("unknown-usd", transaction_date=None, amount="-3", currency="USD"),
    ), ())
    destination = tmp_path / "undated.xlsx"
    XlsxWorkbookWriter().write(snapshot, destination)
    with closing(load_workbook(destination)) as workbook:
        assert list(workbook["月份幣別摘要"].iter_rows(min_row=2, values_only=True)) == [
            (datetime(2026, 9, 1), "TWD", 5, 0, 5, 1),
            ("未知月份", "TWD", 10.25, 0.25, 10, 2),
            ("未知月份", "USD", 0, 3, -3, 1),
        ]
        detail = {row[5]: row for row in workbook["交易明細"].iter_rows(min_row=2, values_only=True)}
        assert set(detail) == {tx.id for tx in snapshot.transactions}
        assert all(detail[tx.id][0] is None for tx in snapshot.transactions if tx.transaction_date is None)


def test_empty_workbook_has_title_headers_filters_widths_and_ownership_marker(tmp_path):
    destination = tmp_path / "empty.xlsx"
    XlsxWorkbookWriter().write(empty_snapshot(), destination)
    with closing(load_workbook(destination)) as workbook:
        assert workbook.properties.title == "家庭收支記錄"
        assert workbook.sheetnames == ["信用卡月支出", "月份幣別摘要", "交易明細", "分類支出", "文件狀態", "_family_finance_hub"]
        for name, last_column in (("信用卡月支出", "K"), ("月份幣別摘要", "F"), ("交易明細", "P"), ("分類支出", "E"), ("文件狀態", "D")):
            sheet = workbook[name]
            assert sheet.max_row == 1
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref == f"A1:{last_column}1"
            assert sheet.sheet_state == "visible"
            assert all(cell.value and cell.font.bold for cell in sheet[1])
            assert all(dimension.width >= 14 for dimension in sheet.column_dimensions.values())
        marker = workbook["_family_finance_hub"]
        assert marker.sheet_state == "veryHidden"
        assert marker["A1"].value == "family-finance-hub:workbook-projection:v1"


def test_utf8_paths_and_chinese_document_states_survive_replacement(tmp_path):
    destination = tmp_path / "家庭資料" / "家庭收支記錄.xlsx"
    tx = transaction()
    active = ExportDocument(tx.document_id, tx.source_filename, "active", 1)
    revoked = ExportDocument("revoked", "已撤銷帳單.csv", "revoked", 17)
    pending = ExportDocument("pending", "待處理帳單.pdf", "pending", 0)
    writer = XlsxWorkbookWriter()
    writer.write(WorkbookSnapshot((tx, transaction("removed")), (active,)), destination)
    writer.write(WorkbookSnapshot((tx,), (pending, revoked, active)), destination)
    with closing(load_workbook(destination)) as workbook:
        documents = {row[0]: row[1:] for row in workbook["文件狀態"].iter_rows(min_row=2, values_only=True)}
        assert documents == {
            active.id: (active.filename, "已入帳", 1),
            revoked.id: (revoked.filename, "已撤銷", 17),
            pending.id: (pending.filename, "尚未建立交易", 0),
        }
        assert workbook["交易明細"].max_row == 2
        assert workbook["交易明細"]["E2"].value == tx.source_filename
        assert workbook["月份幣別摘要"]["F2"].value == 1
        assert workbook["交易明細"].auto_filter.ref == "A1:P2"


@pytest.mark.parametrize("text", [
    '=HYPERLINK("https://example.invalid/", "click")',
    "+SUM(1,2)",
    "-1+2",
    "@SUM(1,2)",
    "\t=1+1",
    "\r=1+1",
    "#N/A",
    "https://example.invalid/",
], ids=["equals", "plus", "minus", "at", "tab", "carriage-return", "error", "url"])
def test_all_untrusted_strings_remain_literals_without_formulas_or_links(tmp_path, text):
    tx = ExportTransaction(text, text, date(2026, 9, 23), text, Decimal("-1"), text, text)
    document = ExportDocument(text, text, text, 1)
    destination = tmp_path / "literal.xlsx"
    XlsxWorkbookWriter().write(WorkbookSnapshot((tx,), (document,)), destination)
    with closing(load_workbook(destination, data_only=False)) as workbook:
        cells = [workbook["交易明細"].cell(2, column) for column in (2, 4, 5, 6, 7)]
        cells.extend(workbook["文件狀態"].cell(2, column) for column in (1, 2, 3))
        cells.append(workbook["月份幣別摘要"]["B2"])
        # XML parsers may normalize literal CR to LF when lxml is not installed.
        assert all(cell.value in (text, text.replace("\r", "\n")) and cell.data_type == "s" for cell in cells)
        assert all(cell.hyperlink is None for cell in cells)
    with ZipFile(destination) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                sheet = ElementTree.fromstring(archive.read(name))
                assert sheet.findall(".//{*}f") == []
                assert sheet.findall(".//{*}hyperlink") == []


@pytest.mark.parametrize("existing_file", [False, True])
def test_failed_atomic_replace_preserves_existing_file_and_cleans_temp(tmp_path, monkeypatch, existing_file):
    destination = tmp_path / "locked.xlsx"
    writer = XlsxWorkbookWriter()
    if existing_file:
        writer.write(empty_snapshot(), destination)
    original = destination.read_bytes() if existing_file else None
    calls = []

    def fail_replace(source, target):
        source = Path(source)
        assert source.parent == destination.parent
        assert target == destination
        assert source.exists()
        calls.append(source)
        with source.open("rb") as stream, closing(load_workbook(stream)) as workbook:
            assert workbook["交易明細"]["F2"].value == "new"
        raise PermissionError("private filename and transaction content")

    monkeypatch.setattr(xlsx.os, "replace", fail_replace)
    with pytest.raises(PermissionError) as error:
        writer.write(WorkbookSnapshot((transaction("new"),), ()), destination)
    assert len(calls) == 1
    assert (destination.read_bytes() if destination.exists() else None) == original
    assert list(tmp_path.iterdir()) == ([destination] if existing_file else [])
    assert "private filename" not in "".join(traceback.format_exception(error.value))


def test_failed_save_cleans_partial_temp_and_preserves_existing_file(tmp_path, monkeypatch):
    destination = tmp_path / "existing.xlsx"
    writer = XlsxWorkbookWriter()
    writer.write(empty_snapshot(), destination)
    original = destination.read_bytes()

    def fail_save(self, stream):
        stream.write(b"partial workbook")
        raise OSError("sensitive source data")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError) as error:
        writer.write(WorkbookSnapshot((transaction(),), ()), destination)
    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]
    assert "sensitive source data" not in "".join(traceback.format_exception(error.value))


@pytest.mark.parametrize("kind", ["ordinary", "wrong-marker", "visible-marker", "corrupt", "empty"])
def test_unowned_workbooks_are_rejected_without_altering_files(tmp_path, kind):
    destination = tmp_path / "personal.xlsx"
    if kind in ("corrupt", "empty"):
        destination.write_bytes(b"private file content" if kind == "corrupt" else b"")
    else:
        with closing(Workbook()) as workbook:
            workbook.active["A1"] = "private content"
            if kind != "ordinary":
                marker = workbook.create_sheet("_family_finance_hub")
                marker["A1"] = (
                    "wrong-owner" if kind == "wrong-marker" else "family-finance-hub:workbook-projection:v1"
                )
                marker.sheet_state = "veryHidden" if kind == "wrong-marker" else "visible"
            workbook.save(destination)
    original = destination.read_bytes()
    with pytest.raises(WorkbookOwnershipError) as error:
        XlsxWorkbookWriter().write(empty_snapshot(), destination)
    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]
    formatted = "".join(traceback.format_exception(error.value))
    assert "private content" not in formatted
    assert str(destination) not in str(error.value)


def test_ownership_is_rechecked_after_rendering(tmp_path, monkeypatch):
    destination = tmp_path / "appeared.xlsx"
    original_save = Workbook.save
    user_content = b"another program created this file"

    def save_with_new_destination(workbook, stream):
        original_save(workbook, stream)
        destination.write_bytes(user_content)

    monkeypatch.setattr(Workbook, "save", save_with_new_destination)
    with pytest.raises(WorkbookOwnershipError):
        XlsxWorkbookWriter().write(empty_snapshot(), destination)
    assert destination.read_bytes() == user_content
    assert list(tmp_path.iterdir()) == [destination]


@pytest.mark.parametrize("invalid_text", ["private\x00content", "x" * 32768], ids=["control", "too-long"])
def test_invalid_text_is_rejected_without_leaking_or_truncating_data(tmp_path, invalid_text):
    destination = tmp_path / "invalid.xlsx"
    tx = replace(transaction(), description=invalid_text)
    with pytest.raises(WorkbookWriteError) as error:
        XlsxWorkbookWriter().write(WorkbookSnapshot((tx,), ()), destination)
    assert invalid_text not in str(error.value)
    assert "private" not in "".join(traceback.format_exception(error.value))
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("amount", ["NaN", "Infinity", "-Infinity"])
def test_nonfinite_amounts_fail_without_silently_producing_empty_cells(tmp_path, amount):
    destination = tmp_path / "invalid.xlsx"
    with pytest.raises(WorkbookWriteError):
        XlsxWorkbookWriter().write(WorkbookSnapshot((transaction(amount=amount),), ()), destination)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.skipif(os.name != "nt", reason="Windows file sharing semantics")
def test_windows_locked_destination_is_preserved_without_temporary_leaks(tmp_path):
    import ctypes
    from ctypes import wintypes

    destination = tmp_path / "locked-by-excel.xlsx"
    writer = XlsxWorkbookWriter()
    writer.write(empty_snapshot(), destination)
    original = destination.read_bytes()
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    create_file.restype = wintypes.HANDLE
    kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
    kernel32.CloseHandle.restype = wintypes.BOOL
    # Permit reads for ownership validation, but deny replacement while the handle is open.
    handle = create_file(str(destination), 0x80000000, 0x00000001, None, 3, 0x80, None)
    if handle == ctypes.c_void_p(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())
    try:
        with pytest.raises(PermissionError):
            writer.write(WorkbookSnapshot((transaction(),), ()), destination)
    finally:
        kernel32.CloseHandle(handle)
    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]
