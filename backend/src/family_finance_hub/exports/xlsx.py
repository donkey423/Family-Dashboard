import os
from contextlib import closing
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .ports import (
    ExportTransaction,
    WorkbookOwnershipError,
    WorkbookSnapshot,
    WorkbookWriteError,
)

_MARKER_SHEET = "_family_finance_hub"
_MARKER_VALUE = "family-finance-hub:workbook-projection:v1"
_OWNERSHIP_ERROR = "Refusing to replace a workbook not owned by this application."
_AMOUNT_FORMAT = "#,##0.00;[Red]-#,##0.00"
_DOCUMENT_STATE_LABELS = {"active": "已入帳", "pending": "尚未建立交易", "revoked": "已撤銷"}


class XlsxWorkbookWriter:
    """Atomically replace only workbooks carrying this projection's marker."""

    def write(self, snapshot: WorkbookSnapshot, destination: Path) -> None:
        temporary_path: Path | None = None
        try:
            _validate_ownership(destination)
            with closing(Workbook()) as workbook:
                _populate_workbook(workbook, snapshot)
                destination.parent.mkdir(parents=True, exist_ok=True)
                with NamedTemporaryFile(
                    mode="w+b",
                    prefix=".family-finance-hub-",
                    suffix=".xlsx.tmp",
                    dir=destination.parent,
                    delete=False,
                ) as temporary:
                    temporary_path = Path(temporary.name)
                    workbook.save(temporary)
                    temporary.flush()
                    os.fsync(temporary.fileno())
            # A destination may have appeared or changed while rendering the snapshot.
            _validate_ownership(destination)
            os.replace(temporary_path, destination)
        except WorkbookWriteError:
            raise
        except PermissionError:
            raise PermissionError("The generated workbook could not be accessed or replaced.") from None
        except OSError:
            raise OSError("Could not write the generated workbook.") from None
        except Exception:
            raise WorkbookWriteError("Could not write the generated workbook.") from None
        finally:
            if temporary_path is not None:
                try:
                    temporary_path.unlink(missing_ok=True)
                except PermissionError:
                    raise PermissionError("Could not clean up the temporary workbook.") from None
                except OSError:
                    raise OSError("Could not clean up the temporary workbook.") from None


def _validate_ownership(destination: Path) -> None:
    if destination.is_symlink():
        raise WorkbookOwnershipError(_OWNERSHIP_ERROR)
    try:
        existing = destination.open("rb")
    except FileNotFoundError:
        return
    with existing:
        try:
            with closing(load_workbook(existing, read_only=True, keep_links=False)) as workbook:
                if _MARKER_SHEET in workbook.sheetnames:
                    marker = workbook[_MARKER_SHEET]
                    cell = marker["A1"]
                    if (
                        marker.sheet_state == "veryHidden"
                        and cell.data_type == "s"
                        and cell.value == _MARKER_VALUE
                    ):
                        return
        except OSError:
            raise
        except Exception:
            raise WorkbookOwnershipError(_OWNERSHIP_ERROR) from None
    raise WorkbookOwnershipError(_OWNERSHIP_ERROR)


def _append_row(
    sheet: Worksheet, values: Iterable[str | date | Decimal | int | None], *, row: int = 1
) -> None:
    for column, value in enumerate(values, start=1):
        if isinstance(value, str) and len(value) > 32767:
            raise WorkbookWriteError("Text exceeds the workbook cell limit.")
        cell = sheet.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            # Override formula/error inference without changing the original text or IDs.
            cell.data_type = "s"
        elif isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
        elif isinstance(value, Decimal):
            cell.number_format = _AMOUNT_FORMAT


def _format_sheet(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{sheet.max_row}"
    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F6B4F")
        cell.alignment = Alignment(vertical="center")
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.data_type == "s")


def _monthly_totals(
    transactions: tuple[ExportTransaction, ...],
) -> dict[tuple[date | None, str], tuple[Decimal, Decimal, int]]:
    totals: dict[tuple[date | None, str], tuple[Decimal, Decimal, int]] = {}
    for transaction in transactions:
        # Statement rows use card-specific semantics in the dedicated sheet;
        # keep this legacy CSV summary from treating card payments as income.
        if transaction.statement_id is not None:
            continue
        if not transaction.amount.is_finite():
            raise WorkbookWriteError("Workbook amounts must be finite numbers.")
        month = (
            transaction.transaction_date.replace(day=1)
            if transaction.transaction_date is not None else None
        )
        key = (month, transaction.currency)
        income, expenses, count = totals.get(key, (Decimal("0"), Decimal("0"), 0))
        if transaction.amount >= 0:
            income += transaction.amount
        else:
            expenses -= transaction.amount
        totals[key] = (income, expenses, count + 1)
    return totals


def _card_monthly_totals(
    transactions: tuple[ExportTransaction, ...],
) -> dict[tuple[date | None, str, str], tuple[Decimal, Decimal, Decimal, Decimal, Decimal, int]]:
    totals: dict[tuple[date | None, str, str], tuple[Decimal, Decimal, Decimal, Decimal, Decimal, int]] = {}
    for transaction in transactions:
        if transaction.statement_id is None:
            continue
        if not transaction.amount.is_finite():
            raise WorkbookWriteError("Workbook amounts must be finite numbers.")
        source_date = transaction.transaction_date or transaction.posting_date
        month = source_date.replace(day=1) if source_date is not None else None
        key = (month, transaction.statement_account or "", transaction.currency)
        purchases, refunds, fees, interest, payments, count = totals.get(
            key, (Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0"), 0)
        )
        if transaction.transaction_kind == "purchase":
            purchases += -transaction.amount
        elif transaction.transaction_kind == "refund":
            refunds += transaction.amount
        elif transaction.transaction_kind == "fee":
            fees += -transaction.amount
        elif transaction.transaction_kind == "interest":
            interest += -transaction.amount
        elif transaction.transaction_kind == "payment":
            payments += transaction.amount
        else:
            raise WorkbookWriteError("Statement transaction kind is not importable.")
        totals[key] = (purchases, refunds, fees, interest, payments, count + 1)
    return totals


def _populate_workbook(workbook: Workbook, snapshot: WorkbookSnapshot) -> None:
    workbook.properties.title = "家庭收支記錄"
    workbook.properties.creator = "family-finance-hub"
    summary = workbook.active
    summary.title = "月份幣別摘要"
    _append_row(summary, ("月份", "幣別", "收入", "支出", "淨額", "交易筆數"))
    totals = _monthly_totals(snapshot.transactions)
    for row, (month, currency) in enumerate(
        sorted(totals, key=lambda key: (key[0] is None, key[0] or date.min, key[1])), start=2
    ):
        income, expenses, count = totals[(month, currency)]
        _append_row(
            summary, (month or "未知月份", currency, income, expenses, income - expenses, count), row=row
        )
        if month is not None:
            summary.cell(row=row, column=1).number_format = "yyyy-mm"
    _format_sheet(summary, (16, 14, 20, 20, 20, 14))

    detail = workbook.create_sheet("交易明細")
    _append_row(detail, (
        "交易日期", "說明", "金額", "幣別", "來源檔名", "交易 ID", "文件 ID",
        "來源類型", "帳戶", "交易類型", "入帳日期", "帳單起日", "帳單迄日", "帳單列號", "Statement ID",
    ))
    for row, transaction in enumerate(
        sorted(
            snapshot.transactions,
            key=lambda item: (item.transaction_date is None, item.transaction_date or date.min, item.id),
        ),
        start=2,
    ):
        _append_row(detail, (
            transaction.transaction_date,
            transaction.description,
            transaction.amount,
            transaction.currency,
            transaction.source_filename,
            transaction.id,
            transaction.document_id,
            transaction.source_type,
            transaction.statement_account,
            transaction.transaction_kind,
            transaction.posting_date,
            transaction.statement_period_start,
            transaction.statement_period_end,
            transaction.statement_line_index,
            transaction.statement_id,
        ), row=row)
    _format_sheet(detail, (16, 60, 20, 14, 50, 40, 40, 24, 24, 18, 16, 16, 16, 14, 40))

    documents = workbook.create_sheet("文件狀態")
    _append_row(documents, ("文件 ID", "檔名", "狀態", "交易筆數"))
    for row, document in enumerate(
        sorted(snapshot.documents, key=lambda item: (item.filename, item.id)), start=2
    ):
        _append_row(
            documents,
            (
                document.id,
                document.filename,
                _DOCUMENT_STATE_LABELS.get(document.state, document.state),
                document.transaction_count,
            ),
            row=row,
        )
    _format_sheet(documents, (40, 50, 18, 14))

    card_summary = workbook.create_sheet("信用卡月支出", 0)
    _append_row(card_summary, (
        "月份", "帳戶", "幣別", "消費", "退款", "淨消費", "費用", "利息", "含費用支出", "繳款", "交易筆數"
    ))
    card_totals = _card_monthly_totals(snapshot.transactions)
    for row, key in enumerate(sorted(card_totals, key=lambda item: (item[0] is None, item[0] or date.min, item[1], item[2])), start=2):
        month, account, currency = key
        purchases, refunds, fees, interest, payments, count = card_totals[key]
        _append_row(card_summary, (
            month or "未知月份",
            account or "未命名帳戶",
            currency,
            purchases,
            refunds,
            purchases - refunds,
            fees,
            interest,
            purchases + fees + interest - refunds,
            payments,
            count,
        ), row=row)
        if month is not None:
            card_summary.cell(row=row, column=1).number_format = "yyyy-mm"
    _format_sheet(card_summary, (16, 24, 14, 18, 18, 18, 18, 18, 20, 18, 14))

    marker = workbook.create_sheet(_MARKER_SHEET)
    _append_row(marker, (_MARKER_VALUE,))
    marker.sheet_state = "veryHidden"
